# -*- coding: utf-8 -*-
"""P0b: LLM 目录三软库首版导入(源=用户 Excel v10「模型库/推理框架库」sheet)。
- 默认 dry_run:打印统计+异常+样例,并落盘 _llm_catalog_dry_run.json 供逐行核对;
- --apply 才写库(upsert by 唯一键,可重复执行)。
用法: python -X utf8 scripts/import_llm_catalog_v10.py [--apply]
"""
import json
import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import openpyxl

from app.models.base import rules_engine  # noqa: E402  (rules schema 引擎)
from app.models.llm_catalog import LlmModel, InferenceFramework, SceneParam  # noqa: E402

XLSX = Path(r"C:\Users\16891\Downloads\本地大模型部署GPU配置器 v10.xlsx")
CONF = {"高": "high", "中": "mid", "低": "low"}


def num(v, cast=float, default=None):
    try:
        return cast(str(v).strip()) if v is not None and str(v).strip() != "" else default
    except (TypeError, ValueError):
        return default


def read_models(ws):
    rows, anomalies = [], []
    for i, row in enumerate(ws.iter_rows(min_row=3, values_only=True), 3):
        vendor, name = (str(row[0] or "").strip()), (str(row[1] or "").strip())
        if not name:
            continue
        params_b = num(row[2], float)
        if params_b is None:
            anomalies.append(f"r{i} {name}: 参数量缺失,跳过")
            continue
        attn = str(row[6] or "GQA").strip().upper() or "GQA"
        conf_cell = str(row[11] or "").strip()
        conf = "low"
        for zh, en in CONF.items():
            if conf_cell.startswith(zh):
                conf = en
                break
        rows.append(dict(
            vendor=vendor, name=name, params_b=params_b,
            default_bits=num(row[3], int, 4) or 4,
            hidden_dim=num(row[4], int), num_layers=num(row[5], int),
            attn=attn, num_heads=num(row[7], int), kv_heads=num(row[8], int),
            vision_gb=num(row[9], float), note=str(row[10] or "").strip() or None,
            confidence=conf,
        ))
    return rows, anomalies


def read_scenes(ws):
    rows = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        name = str(row[13] or "").strip()
        if not name:
            continue
        rows.append(dict(
            name=name, recommend_ctx=num(row[14], int),
            min_tok_s=num(row[15], float),
            need_vision=str(row[16] or "").strip() == "是",
        ))
    return rows


def read_frameworks(ws):
    rows = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        name = str(row[0] or "").strip()
        if not name:
            continue
        rows.append(dict(
            name=name, default_quant=str(row[1] or "").strip() or None,
            coeff_4bit=num(row[2], float), coeff_8bit=num(row[3], float),
            coeff_16bit=num(row[4], float),
            tensor_parallel=str(row[5] or "").startswith("✅"),
            offload=str(row[6] or "").strip() or None,
            overhead_gb=num(row[7], float),
            kv_cache_note=str(row[8] or "").strip() or None,
            note=str(row[9] or "").strip() or None,
        ))
    return rows


def main(apply: bool):
    wb = openpyxl.load_workbook(str(XLSX), data_only=True, read_only=True)
    models, anomalies = read_models(wb["模型库"])
    scenes = read_scenes(wb["模型库"])
    frameworks = read_frameworks(wb["推理框架库"])

    vendors: dict = {}
    for m in models:
        vendors.setdefault(m["vendor"], []).append(m["name"])
    print(f"模型 {len(models)} 个({len(vendors)} 家): " +
          ", ".join(f"{v}×{len(ns)}" for v, ns in vendors.items()))
    print(f"场景 {len(scenes)} 个: {[s['name'] for s in scenes]}")
    print(f"框架 {len(frameworks)} 个: {[f['name'] for f in frameworks]}")
    attn_dist: dict = {}
    for m in models:
        attn_dist[m["attn"]] = attn_dist.get(m["attn"], 0) + 1
    print(f"注意力分布: {attn_dist};置信度分布: " +
          str({c: sum(1 for m in models if m['confidence'] == c) for c in ('high', 'mid', 'low')}))
    for a in anomalies:
        print("  [异常]", a)

    report = Path(__file__).parent / "_llm_catalog_dry_run.json"
    report.write_text(json.dumps(
        {"models": models, "scenes": scenes, "frameworks": frameworks, "anomalies": anomalies},
        ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"明细已落盘 → {report.name}")

    if not apply:
        print("(dry_run 结束,未写库;核对后加 --apply)")
        return

    from sqlalchemy.orm import Session
    from app.models.base import Base
    # 表可能尚未经 startup create_all 建（后端未重启）——幂等补建这三张
    Base.metadata.create_all(rules_engine, tables=[
        LlmModel.__table__, InferenceFramework.__table__, SceneParam.__table__], checkfirst=True)
    with Session(rules_engine) as s:
        for m in models:
            row = s.query(LlmModel).filter_by(vendor=m["vendor"], name=m["name"]).first()
            if row:
                for k, v in m.items():
                    setattr(row, k, v)
            else:
                s.add(LlmModel(**m, created_at=str(date.today())))
        for f in frameworks:
            row = s.query(InferenceFramework).filter_by(name=f["name"]).first()
            if row:
                for k, v in f.items():
                    setattr(row, k, v)
            else:
                s.add(InferenceFramework(**f))
        for i, sc in enumerate(scenes):
            row = s.query(SceneParam).filter_by(name=sc["name"]).first()
            sc["sort_order"] = i
            if row:
                for k, v in sc.items():
                    setattr(row, k, v)
            else:
                s.add(SceneParam(**sc))
        s.commit()
        print(f"已入库: 模型 {len(models)},框架 {len(frameworks)},场景 {len(scenes)}(upsert,可重复执行)")


if __name__ == "__main__":
    main(apply="--apply" in sys.argv)
